"""Gantt plein (barre continue) pour insertion Word / rapport LaTeX."""
from __future__ import annotations

import shutil
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
STAGE_START = date(2026, 3, 6)
STAGE_END = date(2026, 8, 17)

PHASES = [
    ("Cahier des charges", date(2026, 3, 6), date(2026, 3, 19), "#C4A574", "2 sem."),
    ("Conception UML", date(2026, 3, 20), date(2026, 4, 26), "#5B9BD5", "1 mois 1 sem."),
    ("Réalisation", date(2026, 4, 27), date(2026, 7, 26), "#70AD47", "3 mois"),
    ("Tests", date(2026, 7, 27), date(2026, 8, 17), "#ED7D31", "3 sem."),
]

MONTH_NAMES = {
    3: "Mars",
    4: "Avril",
    5: "Mai",
    6: "Juin",
    7: "Juillet",
    8: "Août",
}


def total_days() -> int:
    return (STAGE_END - STAGE_START).days + 1


def weeks() -> list[tuple[int, date, date]]:
    items = []
    d = STAGE_START
    i = 1
    while d <= STAGE_END:
        wend = min(d + timedelta(days=6), STAGE_END)
        items.append((i, d, wend))
        d = wend + timedelta(days=1)
        i += 1
    return items


def phase_for_date(d: date) -> int | None:
    for i, (_n, p1, p2, _c, _dur) in enumerate(PHASES):
        if p1 <= d <= p2:
            return i
    return None


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    name = "calibrib.ttf" if bold else "calibri.ttf"
    path = Path(r"C:\Windows\Fonts") / name
    if not path.exists():
        path = Path(r"C:\Windows\Fonts\arial.ttf")
    return ImageFont.truetype(str(path), size)


def draw_png() -> None:
    ws = weeks()
    n_weeks = len(ws)
    days = total_days()
    left = 32
    right_m = 32
    px_per_day = 10
    chart_w = days * px_per_day
    bar_h = 128
    header_y = 112
    month_h = 38
    week_h = 32
    week_y = header_y + month_h
    bar_y = week_y + week_h
    table_y = bar_y + bar_h + 36
    row_h = 46
    width = left + chart_w + right_m
    height = table_y + 5 * row_h + 28
    chart_right = left + chart_w

    def x_at(d: date, end_of_day: bool = False) -> float:
        offset = (d - STAGE_START).days + (1 if end_of_day else 0)
        return left + chart_w * offset / days

    img = Image.new("RGB", (width, height), "#ffffff")
    draw = ImageDraw.Draw(img)
    f_title = font(36, True)
    f_sub = font(18)
    f_month = font(17, True)
    f_week = font(14, True)
    f_bar = font(18, True)
    f_table = font(17, True)
    f_small = font(16)

    title = "Diagramme de Gantt du projet"
    bbox = draw.textbbox((0, 0), title, font=f_title)
    draw.text(((width - (bbox[2] - bbox[0])) / 2, 16), title, fill="#1f1f1f", font=f_title)
    sub = "Stage du 6 mars au 17 août 2026"
    bbox = draw.textbbox((0, 0), sub, font=f_sub)
    draw.text(((width - (bbox[2] - bbox[0])) / 2, 62), sub, fill="#555555", font=f_sub)

    d = STAGE_START
    while d <= STAGE_END:
        month = d.month
        if d.month == 12:
            last = date(d.year, 12, 31)
        else:
            last = date(d.year, d.month + 1, 1) - timedelta(days=1)
        last = min(last, STAGE_END)
        x1 = x_at(d)
        x2 = x_at(last, end_of_day=True)
        draw.rectangle([x1, header_y, x2, week_y], fill="#2F4F6F")
        label = MONTH_NAMES[month]
        bb = draw.textbbox((0, 0), label, font=f_month)
        draw.text((x1 + (x2 - x1 - (bb[2] - bb[0])) / 2, header_y + 8), label, fill="white", font=f_month)
        d = last + timedelta(days=1)

    for k, (_n, w1, w2) in enumerate(ws):
        x1 = x_at(w1)
        x2 = x_at(w2, end_of_day=True)
        draw.rectangle([x1, week_y, x2, bar_y], fill="#3E5E7E")
        label = f"S{k + 1}"
        bb = draw.textbbox((0, 0), label, font=f_week)
        if (x2 - x1) >= (bb[2] - bb[0] + 4):
            draw.text((x1 + (x2 - x1 - (bb[2] - bb[0])) / 2, week_y + 6), label, fill="white", font=f_week)

    # Continuous phase bar, day by day (handles mid-week transitions).
    d = STAGE_START
    while d <= STAGE_END:
        idx = phase_for_date(d)
        color = PHASES[idx][3] if idx is not None else "#CCCCCC"
        x1 = x_at(d)
        x2 = x_at(d, end_of_day=True)
        draw.rectangle([x1, bar_y, x2, bar_y + bar_h], fill=color)
        d += timedelta(days=1)

    for r, (name, p1, p2, color, dur) in enumerate(PHASES):
        x1 = x_at(p1)
        x2 = x_at(p2, end_of_day=True)
        span = x2 - x1
        if name == "Cahier des charges" and span < 220:
            l1 = "CDC"
        else:
            l1 = name
        l2 = f"({dur})"
        bb1 = draw.textbbox((0, 0), l1, font=f_bar)
        bb2 = draw.textbbox((0, 0), l2, font=f_small)
        cx = (x1 + x2) / 2
        draw.text((cx - (bb1[2] - bb1[0]) / 2, bar_y + 36), l1, fill="#1f1f1f", font=f_bar)
        draw.text((cx - (bb2[2] - bb2[0]) / 2, bar_y + 70), l2, fill="#1f1f1f", font=f_small)

    for _n, w1, _w2 in ws:
        x = x_at(w1)
        draw.line([(x, header_y), (x, bar_y + bar_h)], fill="#5C6B78", width=1)
    draw.line([(chart_right, header_y), (chart_right, bar_y + bar_h)], fill="#5C6B78", width=1)
    draw.rectangle([left, header_y, chart_right, bar_y + bar_h], outline="#2F4F6F", width=2)

    headers = ["Phase", "Début", "Fin", "Durée"]
    span = chart_right - left
    col_x = [left, left + span * 0.42, left + span * 0.62, left + span * 0.80]
    draw.rectangle([left, table_y, chart_right, table_y + row_h], fill="#2F4F6F")
    for i, h in enumerate(headers):
        draw.text((col_x[i] + 16, table_y + 12), h, fill="white", font=f_table)

    for r, (name, p1, p2, color, dur) in enumerate(PHASES):
        y = table_y + (r + 1) * row_h
        draw.rectangle([left, y, chart_right, y + row_h], fill="#F4F6F8" if r % 2 == 0 else "#FFFFFF")
        draw.rectangle([left + 10, y + 12, left + 32, y + 34], fill=color)
        vals = [name, p1.strftime("%d/%m/%Y"), p2.strftime("%d/%m/%Y"), dur]
        for i, v in enumerate(vals):
            draw.text((col_x[i] + (42 if i == 0 else 16), y + 12), v, fill="#1f1f1f", font=f_small)

    for r in range(6):
        y = table_y + r * row_h
        draw.line([(left, y), (chart_right, y)], fill="#8A96A3", width=1)
    draw.rectangle([left, table_y, chart_right, table_y + 5 * row_h], outline="#5C6B78", width=2)

    out = ROOT / "diagramme-gantt.png"
    img.save(out, "PNG", dpi=(300, 300))
    print("Wrote", out, img.size)

    copies = [
        ROOT / "rapport" / "images" / "gantt.png",
        ROOT / "soutenance" / "assets" / "gantt.png",
    ]
    for dest in copies:
        dest.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(out, dest)
        print("Copied", dest)


def draw_xlsx() -> None:
    ws_weeks = weeks()
    wb = Workbook()
    ws = wb.active
    ws.title = "Gantt"

    thin = Border(
        left=Side(style="thin", color="9AA7B2"),
        right=Side(style="thin", color="9AA7B2"),
        top=Side(style="thin", color="9AA7B2"),
        bottom=Side(style="thin", color="9AA7B2"),
    )
    header_fill = PatternFill("solid", fgColor="2F4F6F")
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    n = len(ws_weeks)
    last = get_column_letter(n)
    ws.merge_cells(f"A1:{last}1")
    ws["A1"] = "Diagramme de Gantt du projet — Stage du 6 mars au 17 août 2026"
    ws["A1"].font = Font(name="Calibri", bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24

    for k, (_n, w1, w2) in enumerate(ws_weeks, start=1):
        # Month of the majority of days in the week.
        counts: dict[int, int] = {}
        d = w1
        while d <= w2:
            counts[d.month] = counts.get(d.month, 0) + 1
            d += timedelta(days=1)
        month = max(counts, key=counts.get)
        cell = ws.cell(2, k, MONTH_NAMES[month])
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = thin
        wk = ws.cell(3, k, f"S{k}")
        wk.fill = PatternFill("solid", fgColor="3E5E7E")
        wk.font = header_font
        wk.alignment = center
        wk.border = thin
        ws.column_dimensions[get_column_letter(k)].width = 7

    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 18
    ws.row_dimensions[4].height = 36
    fills = [PatternFill("solid", fgColor=p[3][1:]) for p in PHASES]

    for k, (_n, w1, w2) in enumerate(ws_weeks, start=1):
        mid = w1 + timedelta(days=min(3, (w2 - w1).days))
        idx = phase_for_date(mid)
        c = ws.cell(4, k, "")
        c.border = thin
        c.alignment = center
        c.font = Font(name="Calibri", bold=True, size=10)
        if idx is not None:
            c.fill = fills[idx]
            first = next(
                i
                for i, w in enumerate(ws_weeks, start=1)
                if phase_for_date(w[1] + timedelta(days=min(3, (w[2] - w[1]).days))) == idx
            )
            if k == first:
                c.value = f"{PHASES[idx][0]} ({PHASES[idx][4]})"

    ws["A6"] = "Phase"
    ws["B6"] = "Début"
    ws["C6"] = "Fin"
    ws["D6"] = "Durée"
    for col in range(1, 5):
        cell = ws.cell(6, col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = thin
    for r, (name, p1, p2, color, dur) in enumerate(PHASES):
        row = 7 + r
        vals = [name, p1.strftime("%d/%m/%Y"), p2.strftime("%d/%m/%Y"), dur]
        for col, val in enumerate(vals, start=1):
            cell = ws.cell(row, col, val)
            cell.border = thin
            cell.alignment = Alignment(vertical="center")
            if col == 1:
                cell.fill = PatternFill("solid", fgColor=color[1:])
                cell.font = Font(name="Calibri", bold=True, size=11)
            else:
                cell.font = Font(name="Calibri", size=11)
        ws.row_dimensions[row].height = 22

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1

    out = ROOT / "diagramme-gantt.xlsx"
    wb.save(out)
    print("Wrote", out)


if __name__ == "__main__":
    draw_png()
    draw_xlsx()
