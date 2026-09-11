# -*- coding: utf-8 -*-
"""Génère la présentation de soutenance PFE (widescreen 16:9)."""

from __future__ import annotations

from pathlib import Path
from shutil import copy2

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from PIL import Image

OUT_DIR = Path(__file__).resolve().parent
ASSETS = OUT_DIR / "assets"
LOGO_EMSI = ASSETS / "emsi-logo-crop.png"
LOGO_MEF = ASSETS / "mef-logo.png"
IMG_GANTT = ASSETS / "gantt.png"
IMG_USECASE = ASSETS / "usecase-general.png"
IMG_CLASSE = ASSETS / "classe-general.png"
IMG_SEQ = ASSETS / "seq-repartition.png"
ROOT = OUT_DIR.parent
OUT_FILE = OUT_DIR / "Soutenance-PFE-ETTAHIRI-Mouad.pptx"

NAVY = RGBColor(11, 37, 69)
TEAL = RGBColor(0, 133, 145)
GREEN = RGBColor(34, 112, 54)
INK = RGBColor(31, 41, 55)
MUTED = RGBColor(75, 85, 99)
LINE = RGBColor(209, 213, 219)
WHITE = RGBColor(255, 255, 255)
CREAM = RGBColor(248, 250, 252)
SOFT_RED = RGBColor(153, 27, 27)
BLACK = RGBColor(0, 0, 0)
LIGHT_TEAL = RGBColor(153, 221, 214)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
TOTAL_PAGES = 29

K1 = "Partie 1  ·  Contexte et problématique"
K2 = "Partie 2  ·  Objectifs et démarche"
K3 = "Partie 3  ·  Analyse et conception"
K4 = "Partie 4  ·  Architecture technique"
K5 = "Partie 5  ·  Réalisation"
K6 = "Partie 6  ·  Bilan et perspectives"


def _set_run(run, *, size_pt, bold=False, color=INK, font="Calibri", italic=False):
    run.font.name = font
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def _fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _line(shape, color, width_pt=1.0):
    shape.line.color.rgb = color
    shape.line.width = Pt(width_pt)


def add_textbox(
    slide,
    left,
    top,
    width,
    height,
    text,
    *,
    size=18,
    bold=False,
    italic=False,
    color=INK,
    align=PP_ALIGN.LEFT,
    font="Calibri",
    anchor=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set(
            "anchor",
            {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor],
        )
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    p.clear()
    run = p.add_run()
    run.text = text
    _set_run(run, size_pt=size, bold=bold, italic=italic, color=color, font=font)
    return box


def add_bullets(slide, left, top, width, height, items, *, size=16, color=INK, space_after=10):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(space_after)
        p.clear()
        run = p.add_run()
        run.text = "•   " + item
        _set_run(run, size_pt=size, color=color)
    return box


EMU_PER_INCH = 914400.0


def _to_inches(val):
    """Accepte Inches(), un float en pouces, ou un int EMU issu d'une addition Length."""
    if hasattr(val, "inches"):
        return float(val.inches)
    if isinstance(val, int) and abs(val) > 100:
        return val / EMU_PER_INCH
    return float(val)


def add_picture_fit(slide, path, left, top, max_w, max_h):
    """Insère une image centrée dans un rectangle, sans déformer."""
    im = Image.open(path)
    pw, ph = im.size
    max_w_in = _to_inches(max_w)
    max_h_in = _to_inches(max_h)
    left_in = _to_inches(left)
    top_in = _to_inches(top)
    ratio = pw / ph
    box_ratio = max_w_in / max_h_in
    if ratio > box_ratio:
        w_in, h_in = max_w_in, max_w_in / ratio
    else:
        h_in, w_in = max_h_in, max_h_in * ratio
    x_in = left_in + (max_w_in - w_in) / 2
    y_in = top_in + (max_h_in - h_in) / 2
    slide.shapes.add_picture(str(path), Inches(x_in), Inches(y_in), width=Inches(w_in))


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def add_footer(slide, page, total):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.28), SLIDE_W, Inches(0.22))
    _fill(bar, NAVY)
    add_textbox(
        slide,
        Inches(0.4),
        Inches(7.28),
        Inches(9.8),
        Inches(0.22),
        "Soutenance PFE  ·  14 septembre 2026  ·  ETTAHIRI Mouad  ·  EMSI  ·  MEF",
        size=10,
        color=WHITE,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    add_textbox(
        slide,
        Inches(11.4),
        Inches(7.28),
        Inches(1.5),
        Inches(0.22),
        f"{page} / {total}",
        size=10,
        color=WHITE,
        align=PP_ALIGN.RIGHT,
        anchor=MSO_ANCHOR.MIDDLE,
    )


def add_accent_bar(slide):
    green = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.12), SLIDE_H)
    _fill(green, GREEN)
    teal = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.12), Inches(0), Inches(0.05), SLIDE_H)
    _fill(teal, TEAL)


def _mef_width(height):
    """Largeur du logo MEF (RGB, fond blanc) pour une hauteur donnée."""
    h_in = _to_inches(height)
    with Image.open(LOGO_MEF) as im:
        pw, ph = im.size
    return Inches(h_in * (pw / ph))


def add_header_logos(slide, small=True):
    if small:
        slide.shapes.add_picture(str(LOGO_EMSI), Inches(0.42), Inches(0.18), height=Inches(0.42))
        mef_h = Inches(0.52)
        mef_w = _mef_width(mef_h)
        slide.shapes.add_picture(
            str(LOGO_MEF),
            Inches(13.333) - Inches(0.38) - mef_w,
            Inches(0.14),
            height=mef_h,
        )
    else:
        slide.shapes.add_picture(str(LOGO_EMSI), Inches(0.45), Inches(0.32), height=Inches(0.58))
        mef_h = Inches(0.78)
        mef_w = _mef_width(mef_h)
        slide.shapes.add_picture(
            str(LOGO_MEF),
            Inches(13.333) - Inches(0.42) - mef_w,
            Inches(0.22),
            height=mef_h,
        )


def new_content_slide(prs, title, *, kicker, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), SLIDE_W, SLIDE_H)
    _fill(bg, WHITE)
    add_accent_bar(slide)
    add_header_logos(slide, small=True)
    add_textbox(slide, Inches(0.48), Inches(0.68), Inches(12.2), Inches(0.28), kicker, size=12, bold=True, color=TEAL)
    add_textbox(slide, Inches(0.48), Inches(0.92), Inches(12.2), Inches(0.46), title, size=26, bold=True, color=NAVY)
    add_footer(slide, page, total)
    return slide


def add_cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), SLIDE_W, SLIDE_H)
    _fill(bg, WHITE)

    header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), SLIDE_W, Inches(1.18))
    _fill(header, CREAM)
    add_header_logos(slide, small=False)

    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.18), SLIDE_W, Inches(0.06))
    _fill(line, GREEN)
    line2 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.24), SLIDE_W, Inches(0.025))
    _fill(line2, TEAL)

    add_textbox(
        slide, Inches(0.7), Inches(1.42), Inches(12), Inches(0.28),
        "ROYAUME DU MAROC",
        size=13, bold=True, color=TEAL, align=PP_ALIGN.CENTER,
    )
    add_textbox(
        slide, Inches(0.7), Inches(1.68), Inches(12), Inches(0.40),
        "École Marocaine des Sciences de l'Ingénieur  ·  Campus Rabat",
        size=16, color=NAVY, align=PP_ALIGN.CENTER,
    )

    badge = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.55), Inches(2.16), Inches(6.25), Inches(0.36)
    )
    _fill(badge, NAVY)
    badge.adjustments[0] = 0.5
    add_textbox(
        slide, Inches(3.55), Inches(2.16), Inches(6.25), Inches(0.36),
        "PROJET DE FIN D'ÉTUDES  ·  INGÉNIEUR D'ÉTAT  ·  14 SEPTEMBRE 2026",
        size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
    )

    add_textbox(
        slide, Inches(0.7), Inches(2.58), Inches(12), Inches(0.28),
        "Spécialité Ingénierie Informatique et Réseaux  ·  Option MIAGE",
        size=14, color=MUTED, align=PP_ALIGN.CENTER,
    )
    add_textbox(
        slide, Inches(0.7), Inches(2.90), Inches(11.9), Inches(1.00),
        "Conception et développement d'une application pour\n"
        "l'automatisation de la gestion des concours de recrutement",
        size=22, bold=True, color=NAVY, align=PP_ALIGN.CENTER,
    )
    add_textbox(
        slide, Inches(1.2), Inches(3.92), Inches(10.9), Inches(0.62),
        "Dématérialisation de la gestion des candidatures :\n"
        "import Excel, répartition automatique et convocations",
        size=15, color=INK, align=PP_ALIGN.CENTER,
    )

    cards = [
        ("Réalisé par", "ETTAHIRI Mouad"),
        ("Tuteur de l'école", "Pr. Imane Hilal"),
        ("Tuteur de stage", "M. Tarik Lakhbizi"),
        ("Organisme d'accueil", "Ministère de l'Économie\net des Finances"),
    ]
    left, width, gap = 0.55, 2.95, 0.18
    for i, (label, value) in enumerate(cards):
        x = Inches(left + i * (width + gap))
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(4.78), Inches(width), Inches(1.42))
        _fill(card, CREAM)
        _line(card, LINE, 0.75)
        card.adjustments[0] = 0.12
        add_textbox(
            slide, x + Inches(0.12), Inches(4.88), Inches(width - 0.24), Inches(0.30),
            label, size=11, color=TEAL, align=PP_ALIGN.CENTER, bold=True,
        )
        add_textbox(
            slide, x + Inches(0.12), Inches(5.18), Inches(width - 0.24), Inches(0.88),
            value, size=14, color=NAVY, align=PP_ALIGN.CENTER, bold=True,
        )

    footer = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(6.95), SLIDE_W, Inches(0.55))
    _fill(footer, NAVY)
    add_textbox(
        slide, Inches(0), Inches(6.95), SLIDE_W, Inches(0.55),
        "Soutenance le 14 septembre 2026   ·   Année universitaire 2025 / 2026",
        size=14, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, bold=True,
    )

    add_notes(
        slide,
        "Durée : ~1 min.\n"
        "Bonjour. Je m'appelle Mouad Ettahiri. Je vous présente mon projet de fin d'études, "
        "soutenu le 14 septembre 2026, réalisé à l'EMSI Rabat, option MIAGE, au Ministère de "
        "l'Économie et des Finances, Direction des Affaires administratives et générales. "
        "Le titre officiel : conception et développement d'une application pour "
        "l'automatisation de la gestion des concours de recrutement. "
        "Concrètement : une application web pour importer les candidats, les répartir "
        "automatiquement dans les salles, puis générer et envoyer les convocations.\n"
        "Encadrement : Pr. Imane Hilal (école) et M. Tarik Lakhbizi (stage).",
    )


def add_plan(prs, total_pages):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), SLIDE_W, SLIDE_H)
    _fill(bg, WHITE)
    add_accent_bar(slide)
    add_header_logos(slide, small=True)

    add_textbox(slide, Inches(0.5), Inches(0.72), Inches(12), Inches(0.42), "Plan de la présentation", size=28, bold=True, color=NAVY)
    add_textbox(
        slide, Inches(0.5), Inches(1.14), Inches(12), Inches(0.30),
        "Exposé de 30 minutes  ·  les questions du jury viennent ensuite  ·  14 septembre 2026",
        size=14, color=MUTED,
    )

    parts = [
        ("01", "5 min", "Contexte et problématique", "MEF, constat manuel, question de recherche"),
        ("02", "4 min", "Objectifs et démarche", "Objectifs, cycle en cascade, diagramme de Gantt"),
        ("03", "6 min", "Analyse et conception", "Acteurs, périmètre, cas d'utilisation, classes"),
        ("04", "5 min", "Architecture technique", "Gateway, 6 microservices, JWT, PostgreSQL"),
        ("05", "8 min", "Réalisation", "Import Excel, répartition, convocations, tableau de bord"),
        ("06", "2 min", "Bilan et perspectives", "Apports, limites, évolutions, conclusion"),
    ]
    for i, (num, duree, titre, detail) in enumerate(parts):
        col, row = i % 3, i // 3
        x = Inches(0.45 + col * 4.2)
        y = Inches(1.62 + row * 2.55)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(3.95), Inches(2.32))
        _fill(card, CREAM)
        _line(card, LINE, 0.75)
        card.adjustments[0] = 0.08
        stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(0.12), Inches(2.32))
        _fill(stripe, GREEN if row == 0 else TEAL)
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.28), y + Inches(0.22), Inches(0.55), Inches(0.55))
        _fill(circle, NAVY)
        add_textbox(
            slide, x + Inches(0.28), y + Inches(0.22), Inches(0.55), Inches(0.55),
            num, size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
        )
        pill = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, x + Inches(2.45), y + Inches(0.28), Inches(1.28), Inches(0.38)
        )
        _fill(pill, GREEN if row == 0 else TEAL)
        pill.adjustments[0] = 0.5
        add_textbox(
            slide, x + Inches(2.45), y + Inches(0.28), Inches(1.28), Inches(0.38),
            duree, size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
        )
        add_textbox(slide, x + Inches(0.28), y + Inches(0.95), Inches(3.45), Inches(0.55), titre, size=16, bold=True, color=NAVY)
        add_textbox(slide, x + Inches(0.28), y + Inches(1.48), Inches(3.45), Inches(0.65), detail, size=13, color=MUTED)

    add_footer(slide, 2, total_pages)
    add_notes(
        slide,
        "Durée : ~1 min.\n"
        "Le fil est celui du rapport, condensé pour 30 minutes. "
        "Je commence par le pourquoi (contexte MEF et problématique), "
        "puis le comment on a travaillé (objectifs, cascade, Gantt), "
        "ensuite la conception, l'architecture en microservices, "
        "et le cœur de l'exposé : ce qui a été réellement développé. "
        "Je clos par les apports, les limites et les perspectives.\n"
        "Annonce : « Je reste à votre disposition pour les questions à la fin. »",
    )


def add_organisme(prs, page, total):
    slide = new_content_slide(
        prs,
        "L'organisme d'accueil",
        kicker=K1,
        page=page,
        total=total,
    )

    add_textbox(
        slide, Inches(0.5), Inches(1.42), Inches(7.4), Inches(0.70),
        "Ministère de l'Économie et des Finances",
        size=20, bold=True, color=NAVY,
    )
    add_bullets(
        slide,
        Inches(0.5),
        Inches(2.10),
        Inches(7.5),
        Inches(4.70),
        [
            "Élabore et met en œuvre la politique économique et financière de l'État.",
            "Siège : boulevard Mohammed V, Rabat-Chellah.",
            "Loi de finances, recettes, dépenses, contrôle des finances publiques.",
            "Plus de 17 000 agents dans les directions à réseau, plus de 2 000 en administration centrale.",
            "Modernisation : digitalisation des métiers et gestion du capital humain.",
            "Le stage se déroule à la DAAG, sur l'organisation des concours de recrutement.",
        ],
        size=16,
        space_after=9,
    )

    plate = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.25), Inches(1.55), Inches(4.55), Inches(2.05)
    )
    _fill(plate, WHITE)
    _line(plate, LINE, 0.75)
    plate.adjustments[0] = 0.06
    add_picture_fit(slide, LOGO_MEF, Inches(8.40), Inches(1.62), Inches(4.25), Inches(1.90))

    missions = [
        ("Budget et fiscalité", "Loi de finances, recettes et dépenses"),
        ("Contrôle", "Finances publiques et entreprises publiques"),
        ("RH et SI", "Recrutement, carrière, modernisation"),
    ]
    for i, (t, d) in enumerate(missions):
        y = Inches(3.80 + i * 1.02)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.25), y, Inches(4.55), Inches(0.92))
        _fill(card, CREAM)
        _line(card, LINE, 0.75)
        card.adjustments[0] = 0.12
        add_textbox(slide, Inches(8.42), y + Inches(0.08), Inches(4.22), Inches(0.32), t, size=14, bold=True, color=TEAL)
        add_textbox(slide, Inches(8.42), y + Inches(0.40), Inches(4.22), Inches(0.42), d, size=13, color=INK)

    add_notes(
        slide,
        "Durée : ~1 min 15 s.\n"
        "Le stage se déroule au Ministère de l'Économie et des Finances, à Rabat. "
        "C'est l'administration chargée de la politique économique et financière de l'État : "
        "budget, fiscalité, contrôle des finances publiques. "
        "L'effectif est important, et le Ministère s'est engagé dans la digitalisation de ses métiers, "
        "y compris la gestion des ressources humaines. Mon projet s'inscrit dans cette modernisation : "
        "dématérialiser l'organisation des concours de recrutement.",
    )


def add_daag(prs, page, total):
    slide = new_content_slide(
        prs,
        "La direction d'accueil : la DAAG",
        kicker=K1,
        page=page,
        total=total,
    )
    add_textbox(
        slide, Inches(0.5), Inches(1.42), Inches(12.3), Inches(0.45),
        "Direction des Affaires administratives et générales",
        size=18, color=MUTED,
    )

    blocks = [
        ("01", "Ressources humaines", "Politique RH, déroulement de carrière et action sociale."),
        ("02", "Recrutement", "Coordination des concours et de l'organisation des épreuves."),
        ("03", "Systèmes d'information", "SI et technologies au service des métiers du Ministère."),
    ]
    for i, (num, titre, detail) in enumerate(blocks):
        x = Inches(0.48 + i * 4.20)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.05), Inches(3.95), Inches(2.55))
        _fill(card, CREAM)
        _line(card, LINE, 0.75)
        card.adjustments[0] = 0.08
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.22), Inches(2.22), Inches(0.52), Inches(0.52))
        _fill(circle, NAVY)
        add_textbox(
            slide, x + Inches(0.22), Inches(2.22), Inches(0.52), Inches(0.52),
            num, size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
        )
        add_textbox(slide, x + Inches(0.22), Inches(2.88), Inches(3.50), Inches(0.55), titre, size=17, bold=True, color=NAVY)
        add_textbox(slide, x + Inches(0.22), Inches(3.45), Inches(3.50), Inches(0.95), detail, size=14, color=INK)

    banner = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.48), Inches(4.85), Inches(12.35), Inches(2.05)
    )
    _fill(banner, NAVY)
    banner.adjustments[0] = 0.06
    add_textbox(
        slide, Inches(0.75), Inches(5.00), Inches(11.85), Inches(0.40),
        "Pourquoi la DAAG est le cadre naturel du projet",
        size=16, bold=True, color=LIGHT_TEAL,
    )
    add_textbox(
        slide, Inches(0.75), Inches(5.42), Inches(11.85), Inches(1.25),
        "L'organisation des concours, l'affectation des candidats aux centres et aux salles, "
        "puis l'envoi des convocations relèvent de la gestion des ressources humaines et de la "
        "modernisation des processus. La plateforme vise à assister les gestionnaires, "
        "en automatisant surtout la répartition.",
        size=16, color=WHITE,
    )
    add_notes(
        slide,
        "Durée : ~1 min.\n"
        "Le stage a été effectué à la DAAG. Cette direction pilote les RH, le recrutement "
        "et les systèmes d'information. L'organisation des concours — affectation aux salles "
        "et convocations — est donc bien son métier. D'où le besoin d'un outil pour assister "
        "les gestionnaires, pas pour remplacer le contrôle humain.",
    )


def add_constat(prs, page, total):
    slide = new_content_slide(
        prs,
        "Le constat : un processus encore manuel",
        kicker=K1,
        page=page,
        total=total,
    )
    add_textbox(
        slide, Inches(0.5), Inches(1.42), Inches(12.3), Inches(0.55),
        "Une fois la liste des candidats établie, il reste à les placer et à les convoquer.",
        size=16, color=MUTED,
    )

    problems = [
        ("Affectation lente", "Tableurs, fichiers échangés, risque de dépassement de capacité ou de double affectation."),
        ("Inéquité géographique", "Pas de logique unique : un candidat peut être envoyé loin de sa ville de résidence."),
        ("Convocations fragmentées", "Envois individuels, retards, erreur d'adresse, salle ou place manquante."),
        ("Faible visibilité", "Pas de vue consolidée : salles saturées, non affectés, historique des envois."),
    ]
    for i, (titre, detail) in enumerate(problems):
        col, row = i % 2, i // 2
        x = Inches(0.48 + col * 6.35)
        y = Inches(2.10 + row * 2.35)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(6.10), Inches(2.15))
        _fill(card, CREAM)
        _line(card, LINE, 0.75)
        card.adjustments[0] = 0.08
        mark = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(0.14), Inches(2.15))
        _fill(mark, SOFT_RED)
        add_textbox(slide, x + Inches(0.40), y + Inches(0.28), Inches(5.45), Inches(0.50), titre, size=18, bold=True, color=NAVY)
        add_textbox(slide, x + Inches(0.40), y + Inches(0.85), Inches(5.45), Inches(1.05), detail, size=15, color=INK)

    add_notes(
        slide,
        "Durée : ~1 min 30 s.\n"
        "Le concours ne s'arrête pas à la liste des admis. Les gestionnaires doivent encore "
        "placer chaque candidat dans un centre, un établissement et une salle, puis produire "
        "et envoyer les convocations. Aujourd'hui, cela reste souvent manuel : Excel, fichiers, "
        "envois un par un. Quatre conséquences : erreurs d'affectation, centres trop éloignés, "
        "convocations incomplètes, et aucune vision d'ensemble pour le responsable.",
    )


def add_problematique(prs, page, total):
    slide = new_content_slide(
        prs,
        "Problématique",
        kicker=K1,
        page=page,
        total=total,
    )

    quote = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.48), Inches(1.55), Inches(12.35), Inches(2.55)
    )
    _fill(quote, NAVY)
    quote.adjustments[0] = 0.05
    add_textbox(
        slide, Inches(0.85), Inches(1.75), Inches(11.65), Inches(2.15),
        "Comment concevoir une plateforme qui dématérialise la gestion des candidatures, "
        "automatise la répartition dans les salles, fiabilise les convocations, "
        "et offre aux gestionnaires un suivi opérationnel clair ?",
        size=20, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
    )

    axes = [
        ("Centraliser", "Concours, lieux d'examen, candidats importés par Excel."),
        ("Automatiser", "Répartition selon la capacité et la proximité géographique."),
        ("Notifier", "PDF de convocation et envoi groupé par e-mail."),
        ("Piloter", "Tableau de bord : effectifs, salles, envois."),
    ]
    for i, (titre, detail) in enumerate(axes):
        x = Inches(0.48 + i * 3.18)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(4.35), Inches(3.02), Inches(2.55))
        _fill(card, CREAM)
        _line(card, LINE, 0.75)
        card.adjustments[0] = 0.10
        n = slide.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.18), Inches(4.52), Inches(0.42), Inches(0.42))
        _fill(n, GREEN)
        add_textbox(
            slide, x + Inches(0.18), Inches(4.52), Inches(0.42), Inches(0.42),
            str(i + 1), size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
        )
        add_textbox(slide, x + Inches(0.18), Inches(5.08), Inches(2.66), Inches(0.45), titre, size=16, bold=True, color=NAVY)
        add_textbox(slide, x + Inches(0.18), Inches(5.52), Inches(2.66), Inches(1.15), detail, size=13, color=INK)

    add_notes(
        slide,
        "Durée : ~1 min 15 s.\n"
        "La question de recherche, telle que formulée dans le rapport : comment concevoir "
        "une plateforme qui dématérialise le processus, automatise la répartition, "
        "fiabilise les convocations, et donne une vision claire aux gestionnaires. "
        "Quatre axes de réponse, que je détaillerai ensuite : centraliser les données, "
        "automatiser l'affectation, notifier par PDF et e-mail, piloter via un tableau de bord. "
        "Le contrôle humain est conservé : répartition manuelle possible, aperçu avant envoi.",
    )


def add_objectifs(prs, page, total):
    slide = new_content_slide(
        prs,
        "Objectifs du projet",
        kicker=K2,
        page=page,
        total=total,
    )

    banner = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.48), Inches(1.46), Inches(12.35), Inches(1.15)
    )
    _fill(banner, NAVY)
    banner.adjustments[0] = 0.06
    add_textbox(
        slide, Inches(0.72), Inches(1.58), Inches(11.90), Inches(0.28),
        "Objectif général", size=12, bold=True, color=LIGHT_TEAL,
    )
    add_textbox(
        slide, Inches(0.72), Inches(1.88), Inches(11.90), Inches(0.58),
        "Concevoir et réaliser une plateforme qui assiste les gestionnaires du MEF : "
        "répartition automatique des candidats et convocations fiables.",
        size=16, color=WHITE,
    )

    goals = [
        ("01", "Référentiel", "Concours, centres, établissements et salles, avec capacité."),
        ("02", "Candidats", "Import Excel, consultation, mise à jour, suppression, affectation manuelle."),
        ("03", "Répartition", "Affectation selon le concours, la capacité et la proximité (ville / région)."),
        ("04", "Convocations", "PDF (identité, lieu, place, date) et envoi groupé par e-mail."),
        ("05", "Pilotage", "Tableau de bord : effectifs, remplissage, dernière exécution, historique."),
        ("06", "Sécurité", "Rôles administrateur (lecture + comptes) et gestionnaire (pilotage)."),
        ("07", "Architecture", "Six microservices pour isoler les responsabilités et faciliter la maintenance."),
    ]
    for i, (num, titre, detail) in enumerate(goals):
        col, row = i % 4, i // 4
        w = 3.02
        gap = 3.18
        if row == 0:
            x = Inches(0.48 + col * gap)
        else:
            x = Inches(0.48 + gap / 2 + col * gap)
        y = Inches(2.80 + row * 2.12)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(w), Inches(1.95))
        _fill(card, CREAM)
        _line(card, LINE, 0.75)
        card.adjustments[0] = 0.10
        n = slide.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.16), y + Inches(0.16), Inches(0.40), Inches(0.40))
        _fill(n, GREEN)
        add_textbox(
            slide, x + Inches(0.16), y + Inches(0.16), Inches(0.40), Inches(0.40),
            num, size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
        )
        add_textbox(slide, x + Inches(0.64), y + Inches(0.18), Inches(w - 0.85), Inches(0.38), titre, size=15, bold=True, color=NAVY)
        add_textbox(slide, x + Inches(0.18), y + Inches(0.68), Inches(w - 0.36), Inches(1.12), detail, size=13, color=INK)

    add_notes(
        slide,
        "Durée : ~1 min 20 s.\n"
        "L'objectif général : une plateforme pour assister les gestionnaires, pas pour remplacer "
        "leur contrôle. Sept objectifs concrets, dans l'ordre du métier : d'abord le référentiel "
        "(concours et lieux), puis l'import Excel des candidats, ensuite le moteur de répartition, "
        "les convocations PDF et e-mail, le tableau de bord, la distinction des rôles, et enfin "
        "une architecture en microservices pour que chaque brique reste maintenable. "
        "Il n'y a pas d'espace candidat ni d'inscription en ligne : l'outil est destiné aux agents.",
    )


def add_methodo(prs, page, total):
    slide = new_content_slide(
        prs,
        "Démarche : un cycle en cascade",
        kicker=K2,
        page=page,
        total=total,
    )
    add_textbox(
        slide, Inches(0.5), Inches(1.42), Inches(12.3), Inches(0.55),
        "Les besoins étaient figés dans le cahier des charges : conception d'abord, code ensuite. "
        "Travail seul, avec une validation de l'encadrant à chaque livrable.",
        size=15, color=MUTED,
    )

    phases = [
        ("1", "Cahier des charges", "Acteurs, parcours, import, répartition, convocations, rôles."),
        ("2", "Conception UML", "Cas d'utilisation, séquences et classes avant d'écrire le code."),
        ("3", "Réalisation", "Interface, microservices, base de données, moteur de répartition."),
        ("4", "Tests", "Vérifications, corrections demandées par l'encadrant, contrôle d'ensemble."),
    ]
    for i, (num, titre, detail) in enumerate(phases):
        x = Inches(0.48 + i * 3.18)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.10), Inches(3.02), Inches(3.15))
        _fill(card, CREAM)
        _line(card, LINE, 0.75)
        card.adjustments[0] = 0.08
        head = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Inches(2.10), Inches(3.02), Inches(0.70))
        _fill(head, NAVY)
        add_textbox(
            slide, x, Inches(2.10), Inches(3.02), Inches(0.70),
            f"{num}  ·  {titre}",
            size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
        )
        add_textbox(slide, x + Inches(0.18), Inches(2.98), Inches(2.66), Inches(2.05), detail, size=14, color=INK)
        if i < 3:
            arrow = slide.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW,
                Inches(3.38 + i * 3.18),
                Inches(3.45),
                Inches(0.28),
                Inches(0.22),
            )
            _fill(arrow, TEAL)

    roles = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.48), Inches(5.45), Inches(12.35), Inches(1.45)
    )
    _fill(roles, NAVY)
    roles.adjustments[0] = 0.06
    add_textbox(
        slide, Inches(0.72), Inches(5.58), Inches(12.0), Inches(0.32),
        "Répartition des rôles", size=14, bold=True, color=LIGHT_TEAL,
    )
    add_textbox(
        slide, Inches(0.72), Inches(5.92), Inches(12.0), Inches(0.80),
        "Étudiant : analyse, UML, backend, interfaces, répartition, convocations et tests.   ·   "
        "Encadrant : définition du besoin, validation de chaque étape, pilotage.   ·   "
        "Pas de sprints courts : une porte qualité à la fin de chaque phase.",
        size=15, color=WHITE,
    )
    add_notes(
        slide,
        "Durée : ~1 min 20 s.\n"
        "Pourquoi la cascade et pas l'agile ? Le cahier des charges était posé dès le départ, "
        "et j'étais seul : il valait mieux figer la conception UML avant le code. "
        "Quatre phases, les mêmes que le Gantt : CDC, conception, réalisation, tests. "
        "À chaque livrable, je présentais à l'encadrant ; s'il validait, je passais à la suite, "
        "sinon je corrigeais. Ce n'est pas un déploiement en production : la dernière phase, "
        "ce sont les tests et les retours de l'encadrant, pas une mise en service au Ministère.",
    )


def add_gantt(prs, page, total):
    slide = new_content_slide(
        prs,
        "Planification : cinq mois et demi de stage",
        kicker=K2,
        page=page,
        total=total,
    )
    add_textbox(
        slide, Inches(0.5), Inches(1.40), Inches(12.3), Inches(0.32),
        "Du 6 mars au 17 août 2026  ·  enchaînement séquentiel validé avec l'encadrant",
        size=15, color=MUTED,
    )
    if IMG_GANTT.exists():
        add_picture_fit(slide, IMG_GANTT, Inches(0.48), Inches(1.76), Inches(12.35), Inches(5.30))
    add_notes(
        slide,
        "Durée : ~1 min 20 s.\n"
        "Le stage dure cinq mois et demi. Deux semaines de cahier des charges (6–19 mars), "
        "un mois et une semaine de conception UML (20 mars – 26 avril) : rien n'est codé tant que les diagrammes "
        "ne sont pas validés. Puis trois mois de réalisation (27 avril – 26 juillet) : authentification, "
        "concours et lieux, import, répartition, convocations, tableau de bord. "
        "Enfin trois semaines de tests et corrections (27 juillet – 17 août). "
        "C'est volontairement linéaire : on ne commence pas le code pendant la conception.",
    )


def add_acteurs(prs, page, total):
    slide = new_content_slide(
        prs,
        "Les acteurs du système",
        kicker=K3,
        page=page,
        total=total,
    )
    actors = [
        ("Gestionnaire", "Acteur métier", GREEN,
         "Écriture : concours, lieux, candidats, répartition, convocations. Consulte le tableau de bord."),
        ("Administrateur", "Supervision", TEAL,
         "Lecture seule sur le métier. Seul habilité à créer, modifier ou supprimer les comptes gestionnaires."),
        ("Système", "Traitements", NAVY,
         "Calcule la répartition, assemble les convocations PDF et envoie les e-mails."),
        ("Candidat", "Externe", MUTED,
         "Destinataire de la convocation. Il n'a pas de compte et n'utilise pas l'application."),
    ]
    for i, (titre, role, color, detail) in enumerate(actors):
        col, row = i % 2, i // 2
        x = Inches(0.48 + col * 6.35)
        y = Inches(1.50 + row * 2.70)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(6.12), Inches(2.50))
        _fill(card, CREAM)
        _line(card, LINE, 0.75)
        card.adjustments[0] = 0.08
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(0.16), Inches(2.50))
        _fill(bar, color)
        add_textbox(slide, x + Inches(0.42), y + Inches(0.22), Inches(5.45), Inches(0.42), titre, size=20, bold=True, color=NAVY)
        pill = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, x + Inches(0.42), y + Inches(0.70), Inches(2.35), Inches(0.34)
        )
        _fill(pill, color)
        pill.adjustments[0] = 0.5
        add_textbox(
            slide, x + Inches(0.42), y + Inches(0.70), Inches(2.35), Inches(0.34),
            role, size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
        )
        add_textbox(slide, x + Inches(0.42), y + Inches(1.18), Inches(5.45), Inches(1.10), detail, size=15, color=INK)
    add_notes(
        slide,
        "Durée : ~1 min.\n"
        "Quatre acteurs, deux seulement se connectent. Le gestionnaire pilote tout le métier. "
        "L'administrateur voit les mêmes écrans mais sans bouton d'écriture, et il gère les comptes. "
        "Le système, c'est le moteur : répartition, PDF, e-mails. "
        "Le candidat reçoit sa convocation ; il n'a pas d'espace dans l'application. "
        "Toute action métier suppose une authentification.",
    )


def add_perimetre(prs, page, total):
    slide = new_content_slide(
        prs,
        "Périmètre du système",
        kicker=K3,
        page=page,
        total=total,
    )
    left = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.48), Inches(1.48), Inches(6.12), Inches(5.50)
    )
    _fill(left, CREAM)
    _line(left, LINE, 0.75)
    left.adjustments[0] = 0.05
    head_l = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.48), Inches(1.48), Inches(6.12), Inches(0.52))
    _fill(head_l, GREEN)
    add_textbox(
        slide, Inches(0.48), Inches(1.48), Inches(6.12), Inches(0.52),
        "Inclus", size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
    )
    add_bullets(
        slide, Inches(0.72), Inches(2.12), Inches(5.65), Inches(4.60),
        [
            "Authentification et comptes gestionnaires",
            "CRUD concours, centres, établissements, salles",
            "Import Excel, modification, suppression des candidats",
            "Répartition automatique, historique, réinitialisation, export",
            "PDF de convocation, envoi groupé, historique",
            "Tableau de bord et identité visuelle MEF",
        ],
        size=15,
        space_after=8,
    )

    right = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.80), Inches(1.48), Inches(6.05), Inches(5.50)
    )
    _fill(right, CREAM)
    _line(right, LINE, 0.75)
    right.adjustments[0] = 0.05
    head_r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.80), Inches(1.48), Inches(6.05), Inches(0.52))
    _fill(head_r, SOFT_RED)
    add_textbox(
        slide, Inches(6.80), Inches(1.48), Inches(6.05), Inches(0.52),
        "Hors périmètre", size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
    )
    add_bullets(
        slide, Inches(7.04), Inches(2.12), Inches(5.58), Inches(4.60),
        [
            "Inscription en ligne (les listes d'admis arrivent déjà faites)",
            "Espace candidat (pas de téléchargement autonome)",
            "Notes, jurys, correction, publication des résultats",
            "Application mobile native",
            "Paiement, pré-inscription, RH au-delà de l'organisation matérielle",
        ],
        size=15,
        space_after=8,
        color=INK,
    )
    add_notes(
        slide,
        "Durée : ~1 min 10 s.\n"
        "Le périmètre est volontairement borné : on organise l'épreuve, on ne recrute pas de A à Z. "
        "Les listes d'admis existent déjà ; on les importe. "
        "Pas d'espace candidat, pas de notes, pas de mobile. "
        "Si le jury demande pourquoi : le besoin du Ministère, c'est l'affectation et les convocations, "
        "pas un portail d'inscription. C'est aussi pour tenir le délai de trois mois.",
    )


def add_processus(prs, page, total):
    slide = new_content_slide(
        prs,
        "Le processus métier cible",
        kicker=K3,
        page=page,
        total=total,
    )
    steps = [
        ("1", "S'authentifier", "Session JWT, rôle chargé."),
        ("2", "Saisir le référentiel", "Centres, concours, établissements, salles."),
        ("3", "Importer les candidats", "Fichier Excel des admis, corrections possibles."),
        ("4", "Répartir", "Affectation auto : centre, établissement, salle, place."),
        ("5", "Piloter", "Tableau de bord : effectifs, salles, alertes."),
        ("6", "Convoquer", "Aperçu PDF, puis envoi groupé par e-mail."),
    ]
    for i, (num, titre, detail) in enumerate(steps):
        col, row = i % 3, i // 3
        x = Inches(0.48 + col * 4.20)
        y = Inches(1.52 + row * 2.70)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(3.98), Inches(2.45))
        _fill(card, CREAM)
        _line(card, LINE, 0.75)
        card.adjustments[0] = 0.08
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.22), y + Inches(0.22), Inches(0.55), Inches(0.55))
        _fill(circle, NAVY)
        add_textbox(
            slide, x + Inches(0.22), y + Inches(0.22), Inches(0.55), Inches(0.55),
            num, size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
        )
        add_textbox(slide, x + Inches(0.90), y + Inches(0.28), Inches(2.85), Inches(0.48), titre, size=16, bold=True, color=NAVY)
        add_textbox(slide, x + Inches(0.22), y + Inches(1.00), Inches(3.52), Inches(1.20), detail, size=14, color=INK)
        if col < 2:
            arrow = slide.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW, x + Inches(3.88), y + Inches(1.05), Inches(0.28), Inches(0.20)
            )
            _fill(arrow, TEAL)
    add_notes(
        slide,
        "Durée : ~1 min.\n"
        "C'est le fil que le gestionnaire suit dans l'application. "
        "L'ordre n'est pas libre : d'abord les lieux et les concours, ensuite l'import, "
        "ensuite seulement la répartition, et les convocations uniquement pour les candidats "
        "entièrement affectés. Le tableau de bord sert de contrôle entre la répartition et l'envoi.",
    )


def add_usecase(prs, page, total):
    slide = new_content_slide(
        prs,
        "Diagramme de cas d'utilisation général",
        kicker=K3,
        page=page,
        total=total,
    )
    points = [
        ("Deux acteurs humains", "Gestionnaire : écriture métier. Administrateur : lecture + comptes."),
        ("« include » authentifier", "Aucun cas d'utilisation sans session : la sécurité est dans le modèle, pas seulement dans le code."),
        ("Le candidat n'y figure pas", "Il ne se connecte pas. Il reçoit la convocation à l'extérieur du système."),
    ]
    for i, (titre, detail) in enumerate(points):
        y = Inches(1.48 + i * 1.75)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.45), y, Inches(5.35), Inches(1.60))
        _fill(card, CREAM)
        _line(card, LINE, 0.75)
        card.adjustments[0] = 0.10
        add_textbox(slide, Inches(0.65), y + Inches(0.18), Inches(4.95), Inches(0.40), titre, size=15, bold=True, color=NAVY)
        add_textbox(slide, Inches(0.65), y + Inches(0.60), Inches(4.95), Inches(0.82), detail, size=13, color=INK)

    frame = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.95), Inches(1.46), Inches(6.90), Inches(5.52)
    )
    _fill(frame, WHITE)
    _line(frame, LINE, 0.75)
    frame.adjustments[0] = 0.04
    add_picture_fit(slide, IMG_USECASE, Inches(6.05), Inches(1.54), Inches(6.70), Inches(5.36))
    add_notes(
        slide,
        "Durée : ~1 min 20 s.\n"
        "Le diagramme général, produit avec PlantUML et validé avant le code. "
        "Je retiens trois points pour le jury. Un : le gestionnaire a tous les cas d'écriture, "
        "l'administrateur n'a que la lecture et « Gérer les comptes ». "
        "Deux : chaque cas « include » S'authentifier — ce n'est pas décoratif, c'est une règle. "
        "Trois : le candidat n'apparaît pas, conformément au périmètre. "
        "Les consultations (listes, dashboard, répartition) sont partagées.",
    )


def add_classes(prs, page, total):
    slide = new_content_slide(
        prs,
        "Diagramme de classes général",
        kicker=K3,
        page=page,
        total=total,
    )
    add_picture_fit(slide, IMG_CLASSE, Inches(0.40), Inches(1.42), Inches(12.50), Inches(4.72))
    cap = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.48), Inches(6.22), Inches(12.35), Inches(0.88)
    )
    _fill(cap, NAVY)
    cap.adjustments[0] = 0.08
    add_textbox(
        slide, Inches(0.70), Inches(6.32), Inches(12.0), Inches(0.68),
        "Six paquets, un par microservice. Les pointillés sont des identifiants logiques "
        "(id_centre, numero_concours), pas des clés étrangères. La convocation n'est pas stockée : "
        "elle est assemblée à la lecture.",
        size=13, color=WHITE,
    )
    add_notes(
        slide,
        "Durée : ~1 min 30 s.\n"
        "Six paquets, les mêmes que les microservices plus tard : authentification, lieux, concours, "
        "candidats, répartition, convocations. "
        "Les clés métier sont naturelles : numero_inscription, numero_concours. "
        "Le point important pour le jury : pas de clé étrangère entre bases. "
        "Les pointillés, ce sont des identifiants logiques validés par HTTP. "
        "La convocation n'est pas une table de contenu : on assemble le PDF à partir du candidat affecté ; "
        "on ne stocke que l'historique d'envoi. "
        "Une salle n'a qu'un concours ; un candidat sans affectation complète n'a pas de convocation.",
    )


def add_archi_vue(prs, page, total):
    slide = new_content_slide(
        prs,
        "Architecture : une passerelle, six services, six bases",
        kicker=K4,
        page=page,
        total=total,
    )
    # Navigateur
    nav = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.35), Inches(1.42), Inches(6.60), Inches(0.58))
    _fill(nav, GREEN)
    nav.adjustments[0] = 0.15
    add_textbox(
        slide, Inches(3.35), Inches(1.42), Inches(6.60), Inches(0.58),
        "Navigateur  ·  React 19 + TypeScript + Vite  ·  :5173",
        size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
    )
    add_textbox(
        slide, Inches(3.35), Inches(1.98), Inches(6.60), Inches(0.28),
        "JWT dans l'en-tête Authorization  ·  proxy Vite → gateway",
        size=12, color=MUTED, align=PP_ALIGN.CENTER,
    )

    gw = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.55), Inches(2.32), Inches(10.20), Inches(0.70))
    _fill(gw, NAVY)
    gw.adjustments[0] = 0.08
    add_textbox(
        slide, Inches(1.55), Inches(2.32), Inches(10.20), Inches(0.70),
        "API Gateway  ·  Spring Cloud Gateway  ·  :8080    —    routage par préfixe + CORS    —    pas de JWT, pas de base",
        size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
    )

    services = [
        ("auth", ":8081", "PFE_Data"),
        ("candidat", ":8082", "data_candidats"),
        ("concours", ":8083", "data_concours"),
        ("lieux", ":8084", "data_lieux"),
        ("répartition", ":8085", "data_repartition"),
        ("convocation", ":8086", "data_convocations"),
    ]
    for i, (name, port, db) in enumerate(services):
        x = Inches(0.42 + i * 2.14)
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(3.22), Inches(2.02), Inches(1.55))
        _fill(box, TEAL)
        box.adjustments[0] = 0.10
        add_textbox(slide, x, Inches(3.28), Inches(2.02), Inches(0.55), name, size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_textbox(slide, x, Inches(3.78), Inches(2.02), Inches(0.38), port, size=12, color=WHITE, align=PP_ALIGN.CENTER)
        add_textbox(slide, x, Inches(4.14), Inches(2.02), Inches(0.50), "Spring Boot", size=11, color=WHITE, align=PP_ALIGN.CENTER)

        dbbox = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(4.95), Inches(2.02), Inches(0.72))
        _fill(dbbox, CREAM)
        _line(dbbox, LINE, 0.75)
        dbbox.adjustments[0] = 0.12
        add_textbox(slide, x, Inches(4.95), Inches(2.02), Inches(0.72), db, size=11, bold=True, color=NAVY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    add_textbox(
        slide, Inches(0.45), Inches(5.80), Inches(12.4), Inches(1.20),
        "PostgreSQL  ·  une base par service (database-per-service)  ·  schémas Flyway  ·  pas d'Eureka, pas de Feign\n"
        "Nord–sud : le navigateur ne parle qu'à la gateway.   Est–ouest : RestClient direct entre services, jeton JWT relayé.",
        size=14, color=INK, align=PP_ALIGN.CENTER,
    )
    add_notes(
        slide,
        "Durée : ~1 min.\n"
        "Le navigateur n'appelle jamais un microservice en direct. Tout passe par la gateway, "
        "qui route selon le préfixe : /auth, /api/candidats, /api/concours, etc. "
        "La gateway ne vérifie pas le JWT : c'est un reverse proxy. "
        "Chaque service a sa propre base PostgreSQL. "
        "Entre services, on n'emprunte pas la gateway : RestClient sur les ports 8081 à 8086, "
        "avec le même jeton que l'utilisateur. Pas de bus de messages, pas de découverte de services.",
    )


def add_archi_services(prs, page, total):
    slide = new_content_slide(
        prs,
        "Responsabilités des six microservices",
        kicker=K4,
        page=page,
        total=total,
    )
    rows = [
        ("auth-service", "Login, émission JWT, comptes gestionnaires (admin seulement)."),
        ("candidat-service", "Import Excel, mise à jour, suppression, affectation par lot. Pas de POST unitaire."),
        ("concours-service", "CRUD concours et affectation des centres (id_centre + nom dénormalisé)."),
        ("lieux-service", "Hiérarchie centre → établissement → salle (capacité, un concours par salle)."),
        ("repartition-service", "Orchestrateur : lit les trois référentiels, calcule, écrit les places, historise."),
        ("convocation-service", "Assemble le PDF à la volée, envoi Gmail, journal ENVOYE / ECHEC."),
    ]
    for i, (name, detail) in enumerate(rows):
        y = Inches(1.46 + i * 0.90)
        nbox = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.48), y, Inches(3.35), Inches(0.78))
        _fill(nbox, NAVY if i % 2 == 0 else TEAL)
        nbox.adjustments[0] = 0.12
        add_textbox(
            slide, Inches(0.48), y, Inches(3.35), Inches(0.78),
            name, size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
        )
        dbox = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.98), y, Inches(8.85), Inches(0.78))
        _fill(dbox, CREAM)
        _line(dbox, LINE, 0.6)
        dbox.adjustments[0] = 0.08
        add_textbox(
            slide, Inches(4.18), y, Inches(8.50), Inches(0.78),
            detail, size=14, color=INK, anchor=MSO_ANCHOR.MIDDLE,
        )
    add_notes(
        slide,
        "Durée : ~1 min.\n"
        "Deux services méritent une phrase. La répartition n'a presque pas de données métier : "
        "elle orchestre. Elle lit concours, salles, candidats, calcule en mémoire, puis écrit "
        "les affectations chez candidat-service et garde un historique chez elle. "
        "La convocation n'a que le journal d'envoi : le PDF n'est pas une table, il est assemblé "
        "si le candidat a centre + établissement + salle + place. "
        "Lier une salle à un concours met aussi à jour l'affectation du centre côté concours-service.",
    )


def add_archi_jwt(prs, page, total):
    slide = new_content_slide(
        prs,
        "Sécurité : JWT décentralisé et rôles",
        kicker=K4,
        page=page,
        total=total,
    )
    steps = [
        ("1", "POST /auth/login", "Vérification bcrypt, émission d'un JWT HS256 (sub = login, claim role)."),
        ("2", "sessionStorage", "Clé pfe_access_token. La session meurt à la fermeture de l'onglet. Durée 24 h."),
        ("3", "Chaque service valide", "Même secret JWT_SECRET. Pas d'appel à auth-service pour vérifier le jeton."),
        ("4", "Gateway transparente", "Elle transmet Authorization tel quel. Elle n'authentifie pas, elle n'autorise pas."),
    ]
    for i, (num, titre, detail) in enumerate(steps):
        y = Inches(1.46 + i * 0.88)
        c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.50), y + Inches(0.12), Inches(0.48), Inches(0.48))
        _fill(c, GREEN)
        add_textbox(
            slide, Inches(0.50), y + Inches(0.12), Inches(0.48), Inches(0.48),
            num, size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
        )
        add_textbox(slide, Inches(1.15), y, Inches(6.35), Inches(0.38), titre, size=15, bold=True, color=NAVY)
        add_textbox(slide, Inches(1.15), y + Inches(0.36), Inches(6.35), Inches(0.46), detail, size=13, color=INK)

    box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.70), Inches(1.46), Inches(5.15), Inches(3.45)
    )
    _fill(box, CREAM)
    _line(box, LINE, 0.75)
    box.adjustments[0] = 0.06
    add_textbox(slide, Inches(7.90), Inches(1.58), Inches(4.75), Inches(0.38), "Rôles", size=16, bold=True, color=NAVY)
    add_textbox(
        slide, Inches(7.90), Inches(2.05), Inches(4.75), Inches(2.70),
        "GESTIONNAIRE\nLecture et écriture métier, répartition, envoi.\n\n"
        "ADMINISTRATEUR\nLecture métier. Seul à gérer les comptes.\n\n"
        "Pas de jeton de rafraîchissement : un choix simple, une limite assumée.",
        size=14, color=INK,
    )

    foot = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.48), Inches(5.10), Inches(12.35), Inches(1.85)
    )
    _fill(foot, NAVY)
    foot.adjustments[0] = 0.06
    add_textbox(
        slide, Inches(0.72), Inches(5.25), Inches(11.90), Inches(1.55),
        "Appels internes : le JWT de l'utilisateur est relayé. Un 404 aval devient 400 (référence invalide). "
        "Un 401/403 ou une 5xx aval devient 502. Un service injoignable : 503.\n"
        "Pas de révocation centrale : changer le secret impose de redéployer les six services ressource.",
        size=14, color=WHITE,
    )
    add_notes(
        slide,
        "Durée : ~1 min.\n"
        "Le jeton est sans état : chaque service le vérifie avec le même secret HS256, au moins 32 octets. "
        "La gateway ne fait pas de sécurité métier. "
        "L'administrateur est volontairement en lecture seule sur le concours, pour séparer supervision et opération. "
        "Limite à citer si on nous interroge : pas de refresh token, pas de liste noire. "
        "Le front masque les boutons, mais c'est Spring Security qui tranche.",
    )


def add_archi_seq(prs, page, total):
    slide = new_content_slide(
        prs,
        "Scénario clé : la répartition orchestre trois services",
        kicker=K4,
        page=page,
        total=total,
    )
    phases = [
        ("1. Déclenchement", "POST /api/repartition/run — gestionnaire uniquement."),
        ("2. Lectures", "GET concours, GET salles par concours, GET candidats."),
        ("3. Calcul", "Plan en mémoire : proximité régionale + capacité. Alertes si impossible."),
        ("4. Écriture", "PATCH affectations, puis INSERT run / alertes. Non atomique."),
    ]
    for i, (titre, detail) in enumerate(phases):
        y = Inches(1.46 + i * 1.32)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.42), y, Inches(4.85), Inches(1.18))
        _fill(card, CREAM)
        _line(card, LINE, 0.75)
        card.adjustments[0] = 0.10
        add_textbox(slide, Inches(0.58), y + Inches(0.12), Inches(4.55), Inches(0.38), titre, size=14, bold=True, color=NAVY)
        add_textbox(slide, Inches(0.58), y + Inches(0.50), Inches(4.55), Inches(0.55), detail, size=12, color=INK)

    frame = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.40), Inches(1.42), Inches(7.45), Inches(5.55)
    )
    _fill(frame, WHITE)
    _line(frame, LINE, 0.75)
    frame.adjustments[0] = 0.04
    add_picture_fit(slide, IMG_SEQ, Inches(5.50), Inches(1.50), Inches(7.25), Inches(5.38))
    add_notes(
        slide,
        "Durée : ~1 min 20 s.\n"
        "C'est le diagramme de séquence de la répartition, le scénario le plus architectural. "
        "Le service répartition ne possède pas les salles ni les candidats. Il lit, calcule, écrit. "
        "Les salles viennent de lieux filtrées par numero_concours, pas des affectations concours. "
        "Le calcul est régional : on privilégie un centre de la même région s'il en existe un, "
        "sinon la distance GPS. "
        "Limite : si l'écriture des affectations réussit et que l'historique échoue, on peut avoir "
        "des candidats mis à jour sans run complet. C'est assumé, tracé en ECHEC quand c'est possible.",
    )


def add_archi_stack(prs, page, total):
    slide = new_content_slide(
        prs,
        "Pile technique",
        kicker=K4,
        page=page,
        total=total,
    )
    groups = [
        ("Backend", [
            ("Java 17", "java-logo.png"),
            ("Spring Boot 3.4.4", "spring-boot-logo.png"),
            ("Spring Cloud Gateway", "spring-cloud-logo.jpg"),
            ("Spring Security + JWT", "spring-security-logo.jpg"),
            ("PostgreSQL + Flyway", "postgresql-logo.png"),
            ("Maven", "maven-logo.png"),
        ]),
        ("Métier & front", [
            ("Apache POI (Excel)", "apache-poi-logo.png"),
            ("OpenPDF + Spring Mail", "openpdf-logo.png"),
            ("React 19", "react-logo.png"),
            ("TypeScript + Vite 6", "typescript-logo.png"),
            ("Axios", "axios-logo.png"),
            ("jsPDF / docx", "jspdf-logo.png"),
        ]),
    ]
    y0 = 1.48
    for g, (title, items) in enumerate(groups):
        add_textbox(slide, Inches(0.48), Inches(y0 + g * 2.70), Inches(12), Inches(0.32), title, size=15, bold=True, color=TEAL)
        for i, (label, fname) in enumerate(items):
            x = Inches(0.48 + i * 2.12)
            y = Inches(y0 + 0.38 + g * 2.70)
            card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(2.00), Inches(2.12))
            _fill(card, CREAM)
            _line(card, LINE, 0.75)
            card.adjustments[0] = 0.10
            logo = ROOT / "rapport" / "images" / fname
            if logo.exists():
                add_picture_fit(slide, logo, x + Inches(0.28), y + Inches(0.10), Inches(1.44), Inches(1.18))
            else:
                print(f"Logo manquant : {logo}")
            add_textbox(
                slide, x + Inches(0.08), y + Inches(1.32), Inches(1.84), Inches(0.68),
                label, size=11, bold=True, color=NAVY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
            )
    add_notes(
        slide,
        "Durée : ~40 s.\n"
        "Je ne lis pas tous les logos. Je cite le trio : Java 17 / Spring Boot derrière une gateway, "
        "PostgreSQL une base par service, React + TypeScript devant. "
        "POI pour l'Excel, OpenPDF et Gmail pour les convocations, jsPDF et docx pour les exports "
        "côté gestionnaire. Pas de Docker dans le dépôt : lancement par scripts PowerShell.",
    )


SHOTS = ROOT / "rapport" / "images"


def add_capture_slide(prs, title, subtitle, image_name, notes, page, total):
    slide = new_content_slide(
        prs,
        title,
        kicker=K5,
        page=page,
        total=total,
    )
    add_textbox(slide, Inches(0.50), Inches(1.40), Inches(12.3), Inches(0.32), subtitle, size=14, color=MUTED)
    path = SHOTS / image_name
    add_picture_fit(slide, path, Inches(0.45), Inches(1.76), Inches(12.40), Inches(5.28))
    add_notes(slide, notes)
    return slide


def add_real_connexion(prs, page, total):
    add_capture_slide(
        prs,
        "Connexion",
        "Identité visuelle MEF  ·  aucun écran métier sans authentification",
        "screenshot-connexion.png",
        "Durée : ~45 s.\n"
        "Premier écran : logo du Ministère, identifiant et mot de passe. "
        "Tant que la connexion n'a pas abouti, le reste de l'application est inaccessible. "
        "Le jeton part ensuite dans le sessionStorage. Je ne m'attarde pas : le jury a déjà vu le JWT.",
        page,
        total,
    )


def add_real_dashboard(prs, page, total):
    add_capture_slide(
        prs,
        "Tableau de bord opérationnel",
        "Une vue consolidée : effectifs, salles, dernière répartition, convocations",
        "screenshot-dashboard.png",
        "Durée : ~1 min 10 s.\n"
        "C'est la réponse au constat « pas de vision d'ensemble ». "
        "Six indicateurs : 104 candidats, 104 affectés, 6 concours, 4 centres, 720 places, "
        "104 convocations dont 100 envoyées et 4 échecs. "
        "Le taux d'affectation à 100 % signifie que chaque candidat a centre, établissement, salle et place. "
        "En bas : détail par concours et statut de la dernière run. "
        "Les échecs d'envoi restent visibles : on ne cache pas les problèmes.",
        page,
        total,
    )


def add_real_lieux(prs, page, total):
    add_capture_slide(
        prs,
        "Référentiel des lieux d'examen",
        "Centre → établissement → salle  ·  une salle = un concours et une capacité",
        "screenshot-lieux.png",
        "Durée : ~1 min.\n"
        "L'ordre de saisie est celui du seed : d'abord les centres (ici Casablanca, Marrakech, Rabat, Tanger), "
        "puis les établissements, puis les salles. "
        "Casablanca : établissement Ibn Tofail, salles A, B, C à 30 places, chacune liée à un concours différent. "
        "Les concours liés apparaissent en badges. "
        "C'est ce référentiel que le moteur de répartition lira, pas une table Excel parallèle.",
        page,
        total,
    )


def add_real_candidats(prs, page, total):
    add_capture_slide(
        prs,
        "Candidats : import Excel, pas de saisie unitaire",
        "Importer / réinitialiser  ·  filtres  ·  affectation visible sur chaque fiche",
        "screenshot-candidats.png",
        "Durée : ~1 min 20 s.\n"
        "Les 104 candidats viennent d'un fichier Excel. Il n'y a pas de bouton « Nouveau candidat ». "
        "On peut modifier, supprimer, réinitialiser toute la liste, et affecter manuellement si besoin. "
        "Filtres par nom, ville, concours. "
        "La pastille bleue, c'est le résultat de la répartition : centre, établissement, salle, place. "
        "Exemple à citer : un candidat d'Agadir placé à Marrakech — proximité régionale, pas forcément la même ville.",
        page,
        total,
    )


def add_real_repartition(prs, page, total):
    add_capture_slide(
        prs,
        "Répartition automatique : une exécution, une synthèse",
        "Run #17  ·  104 / 104 affectés  ·  0 alerte  ·  export Word et PDF",
        "screenshot-repartition-synthese.png",
        "Durée : ~1 min 30 s.\n"
        "Le gestionnaire clique sur « Commencer la répartition ». Le système calcule, écrit, historise. "
        "Ici l'exécution 17 : 104 candidats, 104 affectés, zéro alerte, taux 100 %. "
        "Le tableau montre ville d'origine et centre d'affectation : Oujda → Tanger, Fès → Rabat, etc. "
        "Si un candidat ne peut pas être placé, une alerte apparaît (capacité, ville inconnue, concours sans salle). "
        "On peut réinitialiser les places sans relancer, et exporter la synthèse. "
        "Rappel : un run n'est pas une transaction unique entre les six bases.",
        page,
        total,
    )


def add_real_convocations(prs, page, total):
    add_capture_slide(
        prs,
        "Convocations : aperçu, PDF, envoi groupé",
        "Une convocation seulement si l'affectation est complète  ·  100 envoyées, 4 échecs tracés",
        "screenshot-convocations.png",
        "Durée : ~1 min 20 s.\n"
        "Seuls les candidats entièrement placés ont une convocation. "
        "104 prêtes, 104 avec e-mail, 100 envois OK, 4 échecs : un échec n'arrête pas le lot. "
        "Chaque carte : centre, salle, place, date d'examen, bouton « Voir le PDF ». "
        "L'envoi passe par Gmail ; sans mot de passe d'application, l'API répond 503. "
        "Export Word / PDF de la liste côté client, distinct du PDF officiel serveur.",
        page,
        total,
    )


def add_real_pdf_admin(prs, page, total):
    slide = new_content_slide(
        prs,
        "PDF officiel et lecture seule administrateur",
        kicker=K5,
        page=page,
        total=total,
    )
    add_textbox(
        slide, Inches(0.48), Inches(1.40), Inches(6.05), Inches(0.32),
        "Convocation serveur (OpenPDF)", size=13, bold=True, color=TEAL,
    )
    add_textbox(
        slide, Inches(6.80), Inches(1.40), Inches(6.05), Inches(0.32),
        "Vue administrateur : pas de boutons d'écriture", size=13, bold=True, color=TEAL,
    )
    add_picture_fit(slide, SHOTS / "screenshot-convocation-pdf.png", Inches(0.42), Inches(1.74), Inches(6.20), Inches(5.28))
    add_picture_fit(slide, SHOTS / "screenshot-admin-candidats.png", Inches(6.72), Inches(1.74), Inches(6.20), Inches(5.28))
    add_notes(
        slide,
        "Durée : ~1 min.\n"
        "À gauche : le PDF réellement joint à l'e-mail. Identité, concours, date, centre, établissement, "
        "salle, place, consigne d'arriver 30 minutes avant. "
        "À droite : le même écran candidats, rôle administrateur. Badge « lecture seule », "
        "plus de boutons Importer / Modifier / Supprimer, mais un onglet Gestionnaires en plus. "
        "L'UI masque ; Spring Security interdit. "
        "Je clos la démo ici : le parcours import → répartition → convocation est bouclé.",
    )


def add_bilan(prs, page, total):
    slide = new_content_slide(
        prs,
        "Bilan : ce qui a été livré",
        kicker=K6,
        page=page,
        total=total,
    )
    livrables = [
        ("Plateforme intégrée", "Import Excel, lieux, concours, répartition, convocations, tableau de bord."),
        ("Contrôle humain conservé", "Affectation manuelle, aperçu PDF avant envoi, admin en lecture seule."),
        ("Architecture maintenable", "Gateway, 6 microservices, une base par service, JWT."),
        ("Conçu puis développé", "UML validé avant le code, cascade, stage de cinq mois et demi à la DAAG."),
    ]
    for i, (titre, detail) in enumerate(livrables):
        col, row = i % 2, i // 2
        x = Inches(0.48 + col * 6.35)
        y = Inches(1.48 + row * 1.85)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(6.12), Inches(1.70))
        _fill(card, CREAM)
        _line(card, LINE, 0.75)
        card.adjustments[0] = 0.10
        n = slide.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.20), y + Inches(0.22), Inches(0.42), Inches(0.42))
        _fill(n, GREEN)
        add_textbox(
            slide, x + Inches(0.20), y + Inches(0.22), Inches(0.42), Inches(0.42),
            str(i + 1), size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
        )
        add_textbox(slide, x + Inches(0.78), y + Inches(0.24), Inches(5.10), Inches(0.40), titre, size=16, bold=True, color=NAVY)
        add_textbox(slide, x + Inches(0.22), y + Inches(0.78), Inches(5.68), Inches(0.72), detail, size=14, color=INK)

    tests = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.48), Inches(5.28), Inches(12.35), Inches(1.70)
    )
    _fill(tests, NAVY)
    tests.adjustments[0] = 0.06
    add_textbox(slide, Inches(0.72), Inches(5.40), Inches(11.90), Inches(0.32), "Vérifications", size=14, bold=True, color=LIGHT_TEAL)
    add_textbox(
        slide, Inches(0.72), Inches(5.74), Inches(11.90), Inches(1.05),
        "Backend : 34 tests unitaires (candidat, concours, lieux, répartition).   "
        "Frontend : compilation TypeScript + bundle Vite.   "
        "Pas de campagne E2E multi-services.   "
        "L'application n'a pas été mise en exploitation au Ministère pendant le stage.",
        size=15, color=WHITE,
    )
    add_notes(
        slide,
        "Durée : ~45 s.\n"
        "Je clos sur le livrable, pas sur une mise en production. "
        "Le gestionnaire a un outil de bout en bout, avec un contrôle humain. "
        "Les tests couvrent surtout l'algorithme et les clients HTTP, pas un scénario bout-en-bout "
        "avec les six services allumés. Je le dis clairement : le Ministère ne l'a pas encore déployé.",
    )


def add_perspectives(prs, page, total):
    slide = new_content_slide(
        prs,
        "Limites assumées et perspectives",
        kicker=K6,
        page=page,
        total=total,
    )
    left = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.48), Inches(1.48), Inches(6.12), Inches(5.48)
    )
    _fill(left, CREAM)
    _line(left, LINE, 0.75)
    left.adjustments[0] = 0.05
    hl = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.48), Inches(1.48), Inches(6.12), Inches(0.52))
    _fill(hl, NAVY)
    add_textbox(
        slide, Inches(0.48), Inches(1.48), Inches(6.12), Inches(0.52),
        "Technique", size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
    )
    add_bullets(
        slide, Inches(0.72), Inches(2.18), Inches(5.65), Inches(4.50),
        [
            "Jeton d'accès seul : ajouter un jeton de rafraîchissement.",
            "Répartition non atomique : un incident peut laisser des affectations partielles.",
            "Envoi via Gmail : passer à une messagerie institutionnelle.",
            "Secret JWT partagé : la rotation impose de redéployer les six services.",
        ],
        size=15,
        space_after=12,
    )

    right = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.80), Inches(1.48), Inches(6.05), Inches(5.48)
    )
    _fill(right, CREAM)
    _line(right, LINE, 0.75)
    right.adjustments[0] = 0.05
    hr = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.80), Inches(1.48), Inches(6.05), Inches(0.52))
    _fill(hr, GREEN)
    add_textbox(
        slide, Inches(6.80), Inches(1.48), Inches(6.05), Inches(0.52),
        "Métier et exploitation", size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
    )
    add_bullets(
        slide, Inches(7.04), Inches(2.18), Inches(5.58), Inches(4.50),
        [
            "Espace candidat : consulter / télécharger sa convocation, sans inscription en ligne.",
            "Mise en exploitation au Ministère : déploiement, tests complémentaires, formation.",
            "Respect des règles de sécurité de l'administration.",
            "Le périmètre actuel reste l'organisation matérielle des épreuves.",
        ],
        size=15,
        space_after=12,
    )
    add_notes(
        slide,
        "Durée : ~50 s.\n"
        "Je nomme les limites avant qu'on me les pose. "
        "Pas de refresh token, répartition non transactionnelle, Gmail au lieu d'une messagerie MEF. "
        "Ensuite les suites naturelles : un petit espace candidat en consultation seulement, "
        "et un vrai déploiement, qui n'entrait pas dans la durée du stage. "
        "Je ne promets pas l'inscription en ligne : ce n'était pas le besoin.",
    )


def add_merci(prs, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), SLIDE_W, SLIDE_H)
    _fill(bg, NAVY)
    header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), SLIDE_W, Inches(1.18))
    _fill(header, CREAM)
    add_header_logos(slide, small=False)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.18), SLIDE_W, Inches(0.06))
    _fill(line, GREEN)

    add_textbox(
        slide, Inches(0.7), Inches(1.85), Inches(12), Inches(0.40),
        "Soutenance PFE  ·  14 septembre 2026",
        size=16, color=LIGHT_TEAL, align=PP_ALIGN.CENTER,
    )
    add_textbox(
        slide, Inches(0.7), Inches(2.45), Inches(12), Inches(1.10),
        "Merci de votre attention",
        size=40, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
    )
    add_textbox(
        slide, Inches(0.7), Inches(3.55), Inches(12), Inches(0.55),
        "Je reste à votre disposition pour vos questions.",
        size=20, color=WHITE, align=PP_ALIGN.CENTER,
    )

    cards = [
        ("Réalisé par", "ETTAHIRI Mouad"),
        ("Tuteur école", "Pr. Imane Hilal"),
        ("Tuteur stage", "M. Tarik Lakhbizi"),
        ("Organisme", "MEF  ·  DAAG"),
    ]
    for i, (label, value) in enumerate(cards):
        x = Inches(0.70 + i * 3.10)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(4.50), Inches(2.90), Inches(1.45))
        _fill(card, RGBColor(18, 52, 86))
        card.adjustments[0] = 0.12
        add_textbox(
            slide, x + Inches(0.10), Inches(4.62), Inches(2.70), Inches(0.35),
            label, size=12, color=TEAL, align=PP_ALIGN.CENTER, bold=True,
        )
        add_textbox(
            slide, x + Inches(0.10), Inches(5.00), Inches(2.70), Inches(0.72),
            value, size=15, color=WHITE, align=PP_ALIGN.CENTER, bold=True,
        )

    add_textbox(
        slide, Inches(0.7), Inches(6.20), Inches(12), Inches(0.45),
        "EMSI Rabat  ·  Ingénierie Informatique et Réseaux  ·  Option MIAGE",
        size=14, color=RGBColor(180, 198, 220), align=PP_ALIGN.CENTER,
    )
    add_footer(slide, page, total)
    add_notes(
        slide,
        "Durée : ~20 s.\n"
        "Merci. Je remercie Pr. Imane Hilal et M. Tarik Lakhbizi, ainsi que la DAAG. "
        "Je m'arrête ici et j'ouvre les questions. "
        "Ne pas rester debout trop longtemps sur cette slide : passer la parole au jury.",
    )


def _save(prs):
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    candidates = [
        OUT_FILE,
        OUT_DIR / "Soutenance-PFE-ETTAHIRI-Mouad-tmp.pptx",
        OUT_DIR / "Soutenance-PFE-ETTAHIRI-Mouad-v2.pptx",
    ]
    last_error = None
    saved = None
    for path in candidates:
        try:
            prs.save(str(path))
            saved = path
            break
        except PermissionError as exc:
            last_error = exc
    if saved is None:
        raise last_error
    root_copy = ROOT / "Soutenance-PFE-ETTAHIRI-Mouad.pptx"
    if saved.resolve() != root_copy.resolve():
        try:
            copy2(saved, root_copy)
        except PermissionError:
            pass
    return saved


def build():
    if not LOGO_EMSI.exists():
        raise SystemExit(f"Logo EMSI introuvable : {LOGO_EMSI}")
    if not LOGO_MEF.exists():
        raise SystemExit(f"Logo MEF introuvable : {LOGO_MEF}")

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    add_cover(prs)
    add_plan(prs, TOTAL_PAGES)
    add_organisme(prs, 3, TOTAL_PAGES)
    add_daag(prs, 4, TOTAL_PAGES)
    add_constat(prs, 5, TOTAL_PAGES)
    add_problematique(prs, 6, TOTAL_PAGES)
    add_objectifs(prs, 7, TOTAL_PAGES)
    add_methodo(prs, 8, TOTAL_PAGES)
    add_gantt(prs, 9, TOTAL_PAGES)
    add_acteurs(prs, 10, TOTAL_PAGES)
    add_perimetre(prs, 11, TOTAL_PAGES)
    add_processus(prs, 12, TOTAL_PAGES)
    add_usecase(prs, 13, TOTAL_PAGES)
    add_classes(prs, 14, TOTAL_PAGES)
    add_archi_vue(prs, 15, TOTAL_PAGES)
    add_archi_services(prs, 16, TOTAL_PAGES)
    add_archi_jwt(prs, 17, TOTAL_PAGES)
    add_archi_seq(prs, 18, TOTAL_PAGES)
    add_archi_stack(prs, 19, TOTAL_PAGES)
    add_real_connexion(prs, 20, TOTAL_PAGES)
    add_real_dashboard(prs, 21, TOTAL_PAGES)
    add_real_lieux(prs, 22, TOTAL_PAGES)
    add_real_candidats(prs, 23, TOTAL_PAGES)
    add_real_repartition(prs, 24, TOTAL_PAGES)
    add_real_convocations(prs, 25, TOTAL_PAGES)
    add_real_pdf_admin(prs, 26, TOTAL_PAGES)
    add_bilan(prs, 27, TOTAL_PAGES)
    add_perspectives(prs, 28, TOTAL_PAGES)
    add_merci(prs, 29, TOTAL_PAGES)

    saved = _save(prs)
    print(f"OK {saved} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    build()
